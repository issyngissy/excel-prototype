"""Read a soil test workbook and append its rows to the master workbook."""

from __future__ import annotations

import hashlib
import json
import logging
import os
import shutil
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .config import Config

log = logging.getLogger(__name__)

SOURCE_COL = "Source File"
IMPORTED_COL = "Imported At"
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
HEADER_SCAN_ROWS = 20


class MasterLockedError(Exception):
    """The master workbook can't be written, usually because it's open in Excel."""


@dataclass
class ImportResult:
    rows_added: int
    new_columns: list[str]
    skipped_duplicate: bool = False


def is_candidate(path: Path, cfg: Config) -> bool:
    """True if this file is an incoming spreadsheet the agent should import."""
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        return False
    if path.name.startswith(("~$", ".")):  # Excel lock files, macOS resource forks
        return False
    resolved = path.resolve()
    if resolved == cfg.master_file.resolve():
        return False
    for folder in (cfg.processed_folder, cfg.failed_folder):
        if folder and folder.resolve() in resolved.parents:
            return False
    return True


def wait_until_stable(path: Path, checks: int = 3, interval: float = 1.0, timeout: float = 120) -> bool:
    """Wait for a file copy/download to finish: size unchanged across several checks."""
    deadline = time.monotonic() + timeout
    last, same = -1, 0
    while time.monotonic() < deadline:
        try:
            size = path.stat().st_size
        except FileNotFoundError:
            return False
        if size > 0 and size == last:
            same += 1
            if same >= checks:
                return True
        else:
            same = 0
        last = size
        time.sleep(interval)
    return False


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def _norm(value) -> str:
    return " ".join(str(value).strip().lower().split()) if value is not None else ""


def _alias_map(cfg: Config) -> dict[str, str]:
    """Normalised incoming header -> canonical master column name."""
    out: dict[str, str] = {}
    for canonical, aliases in cfg.columns.items():
        out[_norm(canonical)] = canonical
        for alias in aliases:
            out[_norm(alias)] = canonical
    return out


def _find_header_row(rows: list[tuple], aliases: dict[str, str]) -> int:
    """Index of the header row: the row matching the most known column names.

    Falls back to the first row with at least two text cells, so files with
    a title or lab letterhead above the table still work.
    """
    best_idx, best_hits = -1, 0
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        hits = sum(1 for c in row if _norm(c) in aliases)
        if hits > best_hits:
            best_idx, best_hits = i, hits
    if best_idx >= 0:
        return best_idx
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        if sum(1 for c in row if isinstance(c, str) and c.strip()) >= 2:
            return i
    raise ValueError("couldn't find a header row in the first 20 rows")


def read_soil_tests(path: Path, cfg: Config) -> list[dict[str, object]]:
    """Return one dict per data row, keyed by canonical master column name."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[cfg.source_sheet] if cfg.source_sheet else wb.worksheets[0]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    if not rows:
        return []
    aliases = _alias_map(cfg)
    header_idx = _find_header_row(rows, aliases)

    headers: list[str | None] = []
    for cell in rows[header_idx]:
        if cell is None or not str(cell).strip():
            headers.append(None)
        else:
            headers.append(aliases.get(_norm(cell), str(cell).strip()))

    records = []
    for row in rows[header_idx + 1 :]:
        record = {h: v for h, v in zip(headers, row) if h and v is not None and str(v).strip() != ""}
        if record:
            records.append(record)
    return records


def _load_ledger(cfg: Config) -> dict:
    try:
        return json.loads(cfg.ledger_file.read_text())
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


def _save_ledger(cfg: Config, ledger: dict) -> None:
    cfg.ledger_file.write_text(json.dumps(ledger, indent=2))


def _open_master(cfg: Config):
    if cfg.master_file.exists():
        wb = load_workbook(cfg.master_file)
        if cfg.master_sheet in wb.sheetnames:
            ws = wb[cfg.master_sheet]
        else:
            ws = wb.create_sheet(cfg.master_sheet)
    else:
        cfg.master_file.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = cfg.master_sheet

    headers = [c.value for c in ws[1]] if ws.max_row >= 1 else []
    headers = [h for h in headers if h is not None] if any(h is not None for h in headers) else []
    if not headers:
        headers = list(cfg.columns.keys()) + [SOURCE_COL, IMPORTED_COL]
        for i, h in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=h)
    return wb, ws, headers


def _save_master(wb, cfg: Config) -> None:
    """Save via a temp file + rename so a crash never leaves a half-written master."""
    tmp = cfg.master_file.with_name(f".{cfg.master_file.name}.tmp")
    wb.save(tmp)
    try:
        if cfg.master_file.exists():
            shutil.copy2(cfg.master_file, cfg.master_file.with_suffix(cfg.master_file.suffix + ".bak"))
        os.replace(tmp, cfg.master_file)
    except PermissionError as e:
        tmp.unlink(missing_ok=True)
        raise MasterLockedError(str(cfg.master_file)) from e


def append_to_master(path: Path, cfg: Config) -> ImportResult:
    digest = file_hash(path)
    ledger = _load_ledger(cfg)
    if digest in ledger:
        log.info("Skipping %s, identical file already imported as %s", path.name, ledger[digest]["file"])
        return ImportResult(0, [], skipped_duplicate=True)

    records = read_soil_tests(path, cfg)
    wb, ws, headers = _open_master(cfg)

    new_columns: list[str] = []
    for record in records:
        for key in record:
            if key not in headers and key not in new_columns:
                new_columns.append(key)
    if new_columns and not cfg.add_new_columns:
        log.warning("Ignoring columns not in master: %s", ", ".join(new_columns))
        new_columns = []
    for name in new_columns:
        headers.append(name)
        ws.cell(row=1, column=len(headers), value=name)

    col_index = {h: i + 1 for i, h in enumerate(headers)}
    now = datetime.now().replace(microsecond=0)
    for record in records:
        r = ws.max_row + 1
        for key, value in record.items():
            if key in col_index:
                ws.cell(row=r, column=col_index[key], value=value)
        if SOURCE_COL in col_index:
            ws.cell(row=r, column=col_index[SOURCE_COL], value=path.name)
        if IMPORTED_COL in col_index:
            ws.cell(row=r, column=col_index[IMPORTED_COL], value=now)

    _save_master(wb, cfg)
    ledger[digest] = {"file": path.name, "rows": len(records), "imported_at": now.isoformat()}
    _save_ledger(cfg, ledger)
    return ImportResult(len(records), new_columns)


def _move(path: Path, folder: Path | None) -> None:
    if not folder:
        return
    folder.mkdir(parents=True, exist_ok=True)
    dest = folder / path.name
    if dest.exists():
        dest = folder / f"{path.stem}_{datetime.now():%Y%m%d-%H%M%S}{path.suffix}"
    shutil.move(str(path), dest)


def process_file(path: Path, cfg: Config, retries: int = 5, retry_delay: float = 10) -> bool:
    """Import one file end to end. Returns True if the file is finished with.

    Returns False when the master is locked (open in Excel), so the file stays
    in the inbox and the periodic rescan picks it up later.
    """
    if not is_candidate(path, cfg) or not path.exists():
        return True
    if not wait_until_stable(path):
        log.warning("%s never finished copying, will retry on next scan", path.name)
        return False

    for attempt in range(1, retries + 1):
        try:
            result = append_to_master(path, cfg)
        except MasterLockedError:
            log.warning("Master is open/locked (attempt %d/%d), close it in Excel", attempt, retries)
            time.sleep(retry_delay)
            continue
        except Exception:
            log.exception("Failed to import %s, moving to failed folder", path.name)
            _move(path, cfg.failed_folder)
            return True

        if not result.skipped_duplicate:
            log.info("Imported %d rows from %s", result.rows_added, path.name)
            if result.new_columns:
                log.info("Added new master columns: %s", ", ".join(result.new_columns))
        _move(path, cfg.processed_folder)
        return True

    log.error("Gave up on %s for now, master still locked", path.name)
    return False
