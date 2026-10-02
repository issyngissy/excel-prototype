import threading
import time
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from soil_agent import importer
from soil_agent.config import Config
from soil_agent.importer import append_to_master, process_file
from soil_agent.watcher import run


@pytest.fixture
def cfg(tmp_path: Path) -> Config:
    inbox = tmp_path / "Inbox"
    inbox.mkdir()
    return Config(
        watch_folder=inbox,
        master_file=tmp_path / "Master.xlsx",
        processed_folder=inbox / "Processed",
        failed_folder=inbox / "Failed",
        columns={
            "Sample ID": ["sample", "sample no"],
            "Date Sampled": ["date"],
            "pH (CaCl2)": ["ph cacl2"],
            "EC (dS/m)": ["ec"],
        },
    )


REAL_WAIT_UNTIL_STABLE = importer.wait_until_stable


@pytest.fixture(autouse=True)
def fast_stability(monkeypatch):
    monkeypatch.setattr(importer, "wait_until_stable", lambda p, **kw: p.exists())


def make_lab_file(path: Path, rows, headers=("Sample No", "Date", "pH CaCl2", "EC", "Zinc"), title=True):
    wb = Workbook()
    ws = wb.active
    if title:
        ws.append(["ACME Soil Lab, Results Report"])
        ws.append([])
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def master_rows(cfg):
    ws = load_workbook(cfg.master_file)[cfg.master_sheet]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_creates_master_and_maps_aliases(cfg):
    src = make_lab_file(cfg.watch_folder / "lab1.xlsx", [("S1", date(2026, 9, 1), 5.4, 0.12, 1.1)])
    result = append_to_master(src, cfg)

    assert result.rows_added == 1
    assert result.new_columns == ["Zinc"]
    rows = master_rows(cfg)
    assert rows[0] == ["Sample ID", "Date Sampled", "pH (CaCl2)", "EC (dS/m)", "Source File", "Imported At", "Zinc"]
    assert rows[1][:5] == ["S1", rows[1][1], 5.4, 0.12, "lab1.xlsx"]
    assert rows[1][6] == 1.1


def test_appends_to_existing_master_and_skips_blank_rows(cfg):
    make_lab_file(cfg.watch_folder / "a.xlsx", [("S1", None, 5.0, 0.1, None)])
    append_to_master(cfg.watch_folder / "a.xlsx", cfg)
    make_lab_file(cfg.watch_folder / "b.xlsx", [("S2", None, 6.0, 0.2, None), (None,) * 5, ("S3", None, 6.5, 0.3, None)])
    append_to_master(cfg.watch_folder / "b.xlsx", cfg)

    ids = [r[0] for r in master_rows(cfg)[1:]]
    assert ids == ["S1", "S2", "S3"]


def test_duplicate_file_is_not_imported_twice(cfg):
    src = make_lab_file(cfg.watch_folder / "a.xlsx", [("S1", None, 5.0, 0.1, None)])
    append_to_master(src, cfg)
    copy = cfg.watch_folder / "a copy.xlsx"
    copy.write_bytes(src.read_bytes())

    assert append_to_master(copy, cfg).skipped_duplicate
    assert len(master_rows(cfg)) == 2


def test_process_file_moves_to_processed_and_failed(cfg):
    good = make_lab_file(cfg.watch_folder / "good.xlsx", [("S1", None, 5.0, 0.1, None)])
    bad = cfg.watch_folder / "bad.xlsx"
    bad.write_text("not a spreadsheet")

    assert process_file(good, cfg)
    assert process_file(bad, cfg)
    assert (cfg.processed_folder / "good.xlsx").exists()
    assert (cfg.failed_folder / "bad.xlsx").exists()
    assert not good.exists() and not bad.exists()


def test_locked_master_leaves_file_in_inbox(cfg, monkeypatch):
    src = make_lab_file(cfg.watch_folder / "a.xlsx", [("S1", None, 5.0, 0.1, None)])

    def locked(*a, **kw):
        raise importer.MasterLockedError("open in Excel")

    monkeypatch.setattr(importer, "_save_master", locked)
    assert process_file(src, cfg, retries=2, retry_delay=0) is False
    assert src.exists()


def test_ignores_lock_files_and_master(cfg):
    assert not importer.is_candidate(cfg.watch_folder / "~$lab.xlsx", cfg)
    assert not importer.is_candidate(cfg.watch_folder / "notes.csv", cfg)
    assert not importer.is_candidate(cfg.master_file, cfg)
    assert importer.is_candidate(cfg.watch_folder / "lab.xlsx", cfg)


def test_watcher_picks_up_new_file(cfg, monkeypatch):
    # Real stability check: the watcher sees the file before it's fully written.
    monkeypatch.setattr(importer, "wait_until_stable", lambda p: REAL_WAIT_UNTIL_STABLE(p, checks=2, interval=0.3))
    stop = threading.Event()
    t = threading.Thread(target=run, args=(cfg, stop), daemon=True)
    t.start()
    time.sleep(0.5)
    make_lab_file(cfg.watch_folder / "dropped.xlsx", [("W1", None, 5.5, 0.2, None)])

    deadline = time.time() + 10
    while time.time() < deadline and not (cfg.processed_folder / "dropped.xlsx").exists():
        time.sleep(0.2)
    stop.set()

    assert (cfg.processed_folder / "dropped.xlsx").exists()
    assert master_rows(cfg)[1][0] == "W1"
